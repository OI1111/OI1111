from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ── palette ────────────────────────────────────────────────────────────────
NAVY    = '1F3864'
LT_BLUE = 'EBF0F7'
GOLD    = 'C8960C'
GRAY    = '7F7F7F'
OFF_WHT = 'F2F2F2'
MID_BLU = 'B0C4DE'

# ── helpers ────────────────────────────────────────────────────────────────

def set_cell_bg(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)

def remove_cell_borders(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for side in ['top', 'left', 'bottom', 'right']:
        b = OxmlElement(f'w:{side}')
        b.set(qn('w:val'), 'none')
        tcBorders.append(b)
    tcPr.append(tcBorders)

def set_tbl_width(tbl):
    """Set table to 100% page width."""
    tblPr = tbl._tbl.find(qn('w:tblPr'))
    if tblPr is None:
        tblPr = OxmlElement('w:tblPr')
        tbl._tbl.insert(0, tblPr)
    tblW = OxmlElement('w:tblW')
    tblW.set(qn('w:w'), '5000')
    tblW.set(qn('w:type'), 'pct')
    tblPr.append(tblW)

def no_space_tbl_borders(tbl):
    """Remove all table-level borders."""
    tblPr = tbl._tbl.find(qn('w:tblPr'))
    if tblPr is None:
        tblPr = OxmlElement('w:tblPr')
        tbl._tbl.insert(0, tblPr)
    tblBorders = OxmlElement('w:tblBorders')
    for side in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        b = OxmlElement(f'w:{side}')
        b.set(qn('w:val'), 'none')
        tblBorders.append(b)
    tblPr.append(tblBorders)

def banner_tbl(doc, bg, rows_content):
    """
    Create a full-width single-column table with bg color.
    rows_content: list of (text, size, bold, italic, color_hex, space_before, space_after)
    Returns the cell so caller can add more paragraphs if needed.
    """
    tbl = doc.add_table(rows=1, cols=1)
    no_space_tbl_borders(tbl)
    set_tbl_width(tbl)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.rows[0].cells[0]
    set_cell_bg(cell, bg)
    remove_cell_borders(cell)
    first = True
    for (text, size, bold, italic, color_hex, sb, sa) in rows_content:
        if first:
            p = cell.paragraphs[0]
            first = False
        else:
            p = cell.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(sb)
        p.paragraph_format.space_after  = Pt(sa)
        r = p.add_run(text)
        r.font.size = Pt(size)
        r.bold   = bold
        r.italic = italic
        if color_hex:
            r.font.color.rgb = RGBColor.from_string(color_hex)
        else:
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    return cell

def dpar(doc, text='', size=11, bold=False, italic=False, color=None,
         align=WD_ALIGN_PARAGRAPH.CENTER, sb=4, sa=4):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(sb)
    p.paragraph_format.space_after  = Pt(sa)
    if text:
        r = p.add_run(text)
        r.font.size = Pt(size)
        r.bold   = bold
        r.italic = italic
        if color:
            r.font.color.rgb = RGBColor.from_string(color)
    return p

def gold_rule(doc, width=55, sb=2, sa=2):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(sb)
    p.paragraph_format.space_after  = Pt(sa)
    r = p.add_run('─' * width)
    r.font.color.rgb = RGBColor.from_string(GOLD)
    r.font.size = Pt(10)

# ── main generator ─────────────────────────────────────────────────────────

def generate_certificate(output_path,
                         student_name='[STUDENT NAME]',
                         completion_date='________________',
                         achievement_pct=85):

    doc = Document()

    # page setup
    sec = doc.sections[0]
    sec.page_width   = Inches(8.5)
    sec.page_height  = Inches(11)
    sec.left_margin  = Inches(1.0)
    sec.right_margin = Inches(1.0)
    sec.top_margin   = Inches(0.5)
    sec.bottom_margin = Inches(0.5)

    # ── 1. TOP BANNER ──────────────────────────────────────────────────────
    banner_tbl(doc, NAVY, [
        ('BAR-SKILLS',               30, True,  False, None,    18, 4),
        ('CRAFT BARTENDER COURSE',   11, True,  False, MID_BLU,  4, 18),
    ])

    dpar(doc, sb=10, sa=2)  # spacer

    # ── 2. CERTIFICATE TITLE ───────────────────────────────────────────────
    dpar(doc, 'CERTIFICATE OF COMPLETION', size=20, bold=True, color=NAVY,
         align=WD_ALIGN_PARAGRAPH.CENTER, sb=8, sa=4)

    gold_rule(doc, width=60, sb=2, sa=14)

    # ── 3. BODY TEXT ───────────────────────────────────────────────────────
    dpar(doc, 'This is to certify that', size=12, italic=True,
         align=WD_ALIGN_PARAGRAPH.CENTER, sb=4, sa=6)

    dpar(doc, student_name, size=28, bold=True, color=NAVY,
         align=WD_ALIGN_PARAGRAPH.CENTER, sb=6, sa=4)

    gold_rule(doc, width=40, sb=2, sa=10)

    dpar(doc, 'has successfully completed the', size=12, italic=True,
         align=WD_ALIGN_PARAGRAPH.CENTER, sb=4, sa=4)

    dpar(doc, 'Bar-Skills Craft Bartender Course', size=16, bold=True, color=NAVY,
         align=WD_ALIGN_PARAGRAPH.CENTER, sb=4, sa=12)

    # ── 4. STATS BAR (light blue) ──────────────────────────────────────────
    banner_tbl(doc, LT_BLUE, [
        ('8 Sessions  |  90 Minutes Per Session  |  12 Hours Total',
         11, True, False, NAVY, 10, 10),
    ])

    dpar(doc, sb=10, sa=2)  # spacer

    # ── 5. ACHIEVEMENT SCORE (navy) ────────────────────────────────────────
    banner_tbl(doc, NAVY, [
        (f'Overall Achievement Score:   {achievement_pct}%',
         14, True, False, None, 10, 10),
    ])

    dpar(doc, sb=10, sa=2)  # spacer

    # ── 6. COMPLETION CRITERIA BOX ─────────────────────────────────────────
    tbl = doc.add_table(rows=1, cols=1)
    tbl.style = 'Table Grid'
    set_tbl_width(tbl)
    cell = tbl.rows[0].cells[0]
    set_cell_bg(cell, OFF_WHT)

    p0 = cell.paragraphs[0]
    p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p0.paragraph_format.space_before = Pt(8)
    p0.paragraph_format.space_after  = Pt(6)
    r0 = p0.add_run('COMPLETION CRITERIA MET')
    r0.bold = True
    r0.font.size = Pt(10)
    r0.font.color.rgb = RGBColor.from_string(NAVY)

    criteria = [
        'All eight session knowledge tests completed',
        'Comprehensive exam passed at 70% or above  (21/30 minimum)',
        'Speed Track: Negroni and Daiquiri executed simultaneously to standard',
        'Capstone build: stated spec matched the build.  Straw-taste called.  Follow-up questions answered',
    ]
    for c in criteria:
        cp = cell.add_paragraph()
        cp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        cp.paragraph_format.space_before = Pt(2)
        cp.paragraph_format.space_after  = Pt(2)
        cp.paragraph_format.left_indent  = Inches(0.35)
        cr = cp.add_run(f'✓  {c}')
        cr.font.size = Pt(9)
        cr.font.color.rgb = RGBColor.from_string('333333')

    cell.add_paragraph().paragraph_format.space_after = Pt(6)

    dpar(doc, sb=18, sa=4)  # spacer

    # ── 7. SIGNATURE BLOCK ─────────────────────────────────────────────────
    sig = doc.add_table(rows=1, cols=2)
    no_space_tbl_borders(sig)
    set_tbl_width(sig)
    sig.alignment = WD_TABLE_ALIGNMENT.CENTER

    for cell_obj in sig.rows[0].cells:
        remove_cell_borders(cell_obj)

    # left: Oleks
    lc = sig.rows[0].cells[0]
    lp1 = lc.paragraphs[0]
    lp1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    lp1.paragraph_format.space_before = Pt(4)
    lp1.add_run('_' * 30).font.size = Pt(10)

    lp2 = lc.add_paragraph()
    lp2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    lp2.paragraph_format.space_before = Pt(4)
    r = lp2.add_run('Oleks Iurchenko')
    r.bold = True; r.font.size = Pt(10)
    r.font.color.rgb = RGBColor.from_string(NAVY)

    lp3 = lc.add_paragraph()
    lp3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r3 = lp3.add_run('Founder, Bar-Skills  |  WSET Level 3')
    r3.font.size = Pt(9)
    r3.font.color.rgb = RGBColor.from_string(GRAY)

    # right: date
    rc = sig.rows[0].cells[1]
    rp1 = rc.paragraphs[0]
    rp1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rp1.paragraph_format.space_before = Pt(4)
    rp1.add_run('_' * 30).font.size = Pt(10)

    rp2 = rc.add_paragraph()
    rp2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rp2.paragraph_format.space_before = Pt(4)
    r2 = rp2.add_run(completion_date if completion_date != '________________' else 'Date of Completion')
    r2.bold = True; r2.font.size = Pt(10)
    r2.font.color.rgb = RGBColor.from_string(NAVY)

    rp3 = rc.add_paragraph()
    rp3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r5 = rp3.add_run('bar-skills.com  |  info@bar-skills.com')
    r5.font.size = Pt(9)
    r5.font.color.rgb = RGBColor.from_string(GRAY)

    dpar(doc, sb=20, sa=4)  # spacer

    # ── 8. FOOTER BANNER ───────────────────────────────────────────────────
    banner_tbl(doc, NAVY, [
        ('Bar-Skills  |  Kingston, Ontario  |  bar-skills.com  |  info@bar-skills.com',
          8, False, False, MID_BLU, 8, 4),
        ('© 2026 Bar-Skills. All rights reserved.',
          7, False, False, '708090', 2, 8),
    ])

    doc.save(output_path)
    print(f'Saved: {output_path}')


generate_certificate(
    '/home/user/OI1111/BarSkills_Certificate_v2.docx',
    student_name='[STUDENT NAME]',
    completion_date='________________',
    achievement_pct=85,
)
