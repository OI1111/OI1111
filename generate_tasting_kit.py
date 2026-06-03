#!/usr/bin/env python3
"""
Bar-Skills Spirit Tasting Kit Generator
Outputs BarSkills_TastingKit_4pos.docx and BarSkills_TastingKit_6pos.docx
"""

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# Brand palette
DARK_BLUE  = "1F3864"
MID_BLUE   = "8AA5C8"
PALE_BLUE  = "C5D5E8"
LIGHT_BLUE = "EBF0F7"
ALT_GREY   = "F2F2F2"
GREY       = "7F7F7F"
WHITE      = "FFFFFF"


# ── Helpers ─────────────────────────────────────────────────────────────────

def rgb(hex_color):
    h = hex_color.lstrip("#")
    return tuple(int(h[i:i+2], 16) for i in (0, 2, 4))


def set_cell_bg(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def set_cell_margins(cell, top=40, bottom=40, left=72, right=72):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement("w:tcMar")
    for edge, val in [("top", top), ("bottom", bottom),
                      ("left", left), ("right", right)]:
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        tcMar.append(el)
    tcPr.append(tcMar)


def set_row_height(row, height_pts, exact=False):
    trPr = row._tr.get_or_add_trPr()
    trH = OxmlElement("w:trHeight")
    trH.set(qn("w:val"), str(int(height_pts * 20)))
    trH.set(qn("w:hRule"), "exact" if exact else "atLeast")
    trPr.append(trH)


def lock_table(table):
    tblPr = table._tbl.find(qn("w:tblPr"))
    if tblPr is None:
        tblPr = OxmlElement("w:tblPr")
        table._tbl.insert(0, tblPr)
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    tblPr.append(lay)


def set_col_widths(table, widths_inches):
    tbl = table._tbl
    grid = tbl.find(qn("w:tblGrid"))
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        tbl.insert(0, grid)
    else:
        for child in list(grid):
            grid.remove(child)
    for w in widths_inches:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(int(w * 1440)))
        grid.append(col)


def write_cell(cell, text, bold=False, size=8.5, color=DARK_BLUE,
               align=WD_ALIGN_PARAGRAPH.LEFT, italic=False):
    para = cell.paragraphs[0]
    para.clear()
    run = para.add_run(text)
    run.font.bold = bold
    run.font.size = Pt(size)
    run.font.italic = italic
    run.font.name = "Calibri"
    r, g, b = rgb(color)
    run.font.color.rgb = RGBColor(r, g, b)
    para.alignment = align
    para.paragraph_format.space_before = Pt(0)
    para.paragraph_format.space_after = Pt(0)


# ── Sheet builder ────────────────────────────────────────────────────────────

def build_sheet(num_positions):
    doc = Document()

    # Remove default empty paragraph
    for p in doc.paragraphs:
        p._element.getparent().remove(p._element)

    # Page: landscape letter
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width  = Inches(11)
    sec.page_height = Inches(8.5)
    sec.left_margin   = Inches(0.45)
    sec.right_margin  = Inches(0.45)
    sec.top_margin    = Inches(0.35)
    sec.bottom_margin = Inches(0.35)

    # Default style
    ns = doc.styles["Normal"]
    ns.font.name = "Calibri"
    ns.paragraph_format.space_before = Pt(0)
    ns.paragraph_format.space_after  = Pt(0)

    # Column geometry
    usable   = 11 - 0.45 - 0.45   # 10.1"
    label_w  = 1.05
    pos_w    = round((usable - label_w) / num_positions, 4)
    col_widths = [label_w] + [pos_w] * num_positions

    # Content row definitions: (label, min_height_pts, label_bg, pos_bg, hint_text)
    ROWS = [
        ("Spirit",         72,  ALT_GREY,   WHITE,  ""),
        ("Distillery",     36,  LIGHT_BLUE, WHITE,  ""),
        ("Appearance",     52,  ALT_GREY,   WHITE,  ""),
        ("Nose",           80,  LIGHT_BLUE, WHITE,  ""),
        ("Palate",         80,  ALT_GREY,   WHITE,  ""),
        ("Finish",         56,  LIGHT_BLUE, WHITE,  ""),
        ("Template Fit",   44,  ALT_GREY,   WHITE,
            "Old Fashioned  /  Sour  /  Highball  /  Stirred"),
        ("Rating",         36,  LIGHT_BLUE, WHITE,  "          / 10"),
        ("Notes",          52,  ALT_GREY,   WHITE,  ""),
    ]

    total_table_rows = 3 + len(ROWS)  # header + meta + col headers + content
    total_cols = 1 + num_positions

    table = doc.add_table(rows=total_table_rows, cols=total_cols)
    table.style = "Table Grid"
    lock_table(table)
    set_col_widths(table, col_widths)

    ri = 0  # row index

    # ── Row 0: Brand header ──────────────────────────────────────────────────
    hrow = table.rows[ri]
    set_row_height(hrow, 32, exact=True)
    hrow.cells[0].merge(hrow.cells[-1])
    hc = hrow.cells[0]
    set_cell_bg(hc, DARK_BLUE)
    set_cell_margins(hc, top=60, bottom=60, left=144, right=144)

    para = hc.paragraphs[0]
    para.clear()
    para.alignment = WD_ALIGN_PARAGRAPH.LEFT

    def add_run(p, text, bold=False, size=10, hex_color=WHITE):
        run = p.add_run(text)
        run.font.name = "Calibri"
        run.font.bold = bold
        run.font.size = Pt(size)
        r, g, b = rgb(hex_color)
        run.font.color.rgb = RGBColor(r, g, b)

    add_run(para, "BAR-SKILLS", bold=True, size=13, hex_color=WHITE)
    add_run(para, "     Spirit Tasting Notes", bold=False, size=9, hex_color=PALE_BLUE)
    add_run(para, f"     {num_positions} positions", bold=False, size=7.5, hex_color=MID_BLUE)

    # Right-side: bar-skills.com
    right_para = hc.add_paragraph()
    right_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    right_para.paragraph_format.space_before = Pt(0)
    right_para.paragraph_format.space_after = Pt(0)
    add_run(right_para, "bar-skills.com", bold=False, size=7.5, hex_color=MID_BLUE)

    ri += 1

    # ── Row 1: Meta info ─────────────────────────────────────────────────────
    mrow = table.rows[ri]
    set_row_height(mrow, 22, exact=True)
    mrow.cells[0].merge(mrow.cells[-1])
    mc = mrow.cells[0]
    set_cell_bg(mc, LIGHT_BLUE)
    set_cell_margins(mc, top=28, bottom=28, left=144, right=144)

    mpara = mc.paragraphs[0]
    mpara.clear()
    mpara.alignment = WD_ALIGN_PARAGRAPH.LEFT

    for label, space in [("Date", 22), ("Event / Location", 26), ("Tasted by", 16)]:
        run = mpara.add_run(f"  {label}: ")
        run.font.bold = True
        run.font.size = Pt(7)
        run.font.name = "Calibri"
        r, g, b = rgb(DARK_BLUE)
        run.font.color.rgb = RGBColor(r, g, b)
        run2 = mpara.add_run("_" * space + "     ")
        run2.font.size = Pt(7)
        run2.font.name = "Calibri"
        r, g, b = rgb(GREY)
        run2.font.color.rgb = RGBColor(r, g, b)

    ri += 1

    # ── Row 2: Column headers ────────────────────────────────────────────────
    crow = table.rows[ri]
    set_row_height(crow, 20, exact=True)

    lc = crow.cells[0]
    set_cell_bg(lc, DARK_BLUE)
    set_cell_margins(lc, top=28, bottom=28, left=72, right=40)
    write_cell(lc, "Position", bold=True, size=7, color=PALE_BLUE,
               align=WD_ALIGN_PARAGRAPH.LEFT)

    for i in range(1, num_positions + 1):
        c = crow.cells[i]
        set_cell_bg(c, DARK_BLUE)
        set_cell_margins(c, top=28, bottom=28, left=60, right=60)
        write_cell(c, f"# {i}", bold=True, size=8, color=WHITE,
                   align=WD_ALIGN_PARAGRAPH.CENTER)

    ri += 1

    # ── Content rows ─────────────────────────────────────────────────────────
    for label, height, label_bg, pos_bg, hint in ROWS:
        robj = table.rows[ri]
        set_row_height(robj, height, exact=False)

        # Label cell
        lc = robj.cells[0]
        set_cell_bg(lc, label_bg)
        set_cell_margins(lc, top=36, bottom=36, left=72, right=36)
        write_cell(lc, label, bold=True, size=7.5, color=DARK_BLUE)

        # Position cells
        for i in range(1, num_positions + 1):
            c = robj.cells[i]
            set_cell_bg(c, pos_bg)
            set_cell_margins(c, top=28, bottom=28, left=60, right=60)
            if hint:
                write_cell(c, hint, size=6.5, color=GREY,
                           align=WD_ALIGN_PARAGRAPH.CENTER, italic=True)
            else:
                c.paragraphs[0].clear()

        ri += 1

    return doc


def main():
    for n in [4, 6]:
        doc = build_sheet(n)
        fname = f"BarSkills_TastingKit_{n}pos.docx"
        doc.save(fname)
        print(f"Saved: {fname}")


if __name__ == "__main__":
    main()
